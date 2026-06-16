import io
from datetime import datetime

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


DAYS = ["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]


def clean_code(value: object) -> str:
    """Normalize shift codes so D/d, n/N, off/OFF all match."""
    if value is None:
        return ""
    return str(value).strip().upper()


def make_resident_names(n_residents: int) -> list[str]:
    default_names = ["Resident 1", "Resident 2", "Resident 3", "Resident 4", "Resident 5", "Resident 6"]
    names = []

    for i in range(n_residents):
        default = default_names[i] if i < len(default_names) else f"Resident {i + 1}"
        names.append(
            st.sidebar.text_input(
                f"Resident {i + 1}",
                value=default,
                key=f"resident_{i}",
            ).strip()
            or default
        )

    return names


def make_shift_definitions() -> pd.DataFrame:
    st.sidebar.subheader("Shift definitions")

    default_defs = pd.DataFrame(
        {
            "Code": ["D", "N", "OFF", "POST"],
            "Hours": [12, 13, 0, 0],
        }
    )

    shift_defs = st.sidebar.data_editor(
        default_defs,
        hide_index=True,
        num_rows="dynamic",
        use_container_width=True,
        column_config={
            "Code": st.column_config.TextColumn("Code", required=True),
            "Hours": st.column_config.NumberColumn(
                "Hours",
                min_value=0,
                step=0.5,
                required=True,
            ),
        },
        key="shift_defs_editor",
    )

    shift_defs = shift_defs.copy()
    shift_defs["Code"] = shift_defs["Code"].map(clean_code)
    shift_defs["Hours"] = pd.to_numeric(shift_defs["Hours"], errors="coerce").fillna(0)
    shift_defs = shift_defs[shift_defs["Code"] != ""]
    shift_defs = shift_defs.drop_duplicates(subset=["Code"], keep="first").reset_index(drop=True)

    if shift_defs.empty:
        st.sidebar.warning("Add at least one shift code. Reverting to defaults.")
        shift_defs = default_defs

    return shift_defs


def build_blank_schedule(residents: list[str], num_weeks: int, default_code: str) -> pd.DataFrame:
    rows = []

    for week in range(1, num_weeks + 1):
        for resident in residents:
            row = {"Week": week, "Resident": resident}
            for day in DAYS:
                row[day] = default_code
            rows.append(row)

    return pd.DataFrame(rows)


def reconcile_schedule(
    existing: pd.DataFrame | None,
    residents: list[str],
    num_weeks: int,
    default_code: str,
) -> pd.DataFrame:
    """Keep prior edits when controls change, but rebuild missing resident/week rows."""
    blank = build_blank_schedule(residents, num_weeks, default_code)

    if existing is None or existing.empty:
        return blank

    existing = existing.copy()

    for col in ["Week", "Resident", *DAYS]:
        if col not in existing.columns:
            existing[col] = default_code if col in DAYS else None

    merged = blank[["Week", "Resident"]].merge(
        existing[["Week", "Resident", *DAYS]],
        on=["Week", "Resident"],
        how="left",
    )

    for day in DAYS:
        merged[day] = merged[day].fillna(default_code).map(clean_code)

    return merged[["Week", "Resident", *DAYS]]


def sync_week_editor(week: int) -> None:
    """
    Save one week's data_editor dropdown edits immediately.

    This prevents the "clicked once but it snapped back" Streamlit behavior.
    """
    key = f"schedule_editor_week_{week}"

    if "schedule_df" not in st.session_state:
        return

    editor_state = st.session_state.get(key, {})
    df = st.session_state.schedule_df.copy()

    week_indices = df.index[df["Week"] == week].tolist()

    # Current Streamlit format:
    # {"edited_rows": {0: {"Mon": "D"}}}
    edited_rows = editor_state.get("edited_rows", {})

    for row_idx, updates in edited_rows.items():
        row_idx = int(row_idx)

        if row_idx >= len(week_indices):
            continue

        actual_idx = week_indices[row_idx]

        for col, value in updates.items():
            if col in DAYS:
                df.at[actual_idx, col] = clean_code(value)

    # Older Streamlit fallback:
    # {"edited_cells": {"0:Mon": "D"}}
    edited_cells = editor_state.get("edited_cells", {})

    for cell_key, value in edited_cells.items():
        try:
            row_text, col = str(cell_key).split(":", 1)
            row_idx = int(row_text)
        except ValueError:
            continue

        if row_idx >= len(week_indices):
            continue

        actual_idx = week_indices[row_idx]

        if col in DAYS:
            df.at[actual_idx, col] = clean_code(value)

    st.session_state.schedule_df = df


def add_week_calculations(week_df: pd.DataFrame, shift_defs: pd.DataFrame) -> pd.DataFrame:
    """Add visible weekly counts next to each week's calendar table."""
    hours_map = dict(zip(shift_defs["Code"], shift_defs["Hours"]))

    out = week_df.copy()

    for day in DAYS:
        out[day] = out[day].map(clean_code)

    out["D Count"] = (out[DAYS] == "D").sum(axis=1)
    out["N Count"] = (out[DAYS] == "N").sum(axis=1)

    out["Worked Shifts"] = out[DAYS].apply(
        lambda row: sum(hours_map.get(clean_code(x), 0) > 0 for x in row),
        axis=1,
    )

    out["Hours"] = out[DAYS].apply(
        lambda row: sum(hours_map.get(clean_code(x), 0) for x in row),
        axis=1,
    )

    out["OFF Days"] = (out[DAYS] == "OFF").sum(axis=1)
    out["POST Days"] = (out[DAYS] == "POST").sum(axis=1)

    return out


def summarize_schedule(
    schedule_df: pd.DataFrame,
    shift_defs: pd.DataFrame,
    residents: list[str],
    num_weeks: int,
) -> pd.DataFrame:
    shift_codes = shift_defs["Code"].tolist()
    hours_map = dict(zip(shift_defs["Code"], shift_defs["Hours"]))

    long_df = schedule_df.melt(
        id_vars=["Week", "Resident"],
        value_vars=DAYS,
        var_name="Day",
        value_name="Code",
    )

    long_df["Code"] = long_df["Code"].map(clean_code)
    long_df["Hours"] = long_df["Code"].map(hours_map).fillna(0)

    code_counts = (
        pd.crosstab(long_df["Resident"], long_df["Code"])
        .reindex(index=residents, columns=shift_codes, fill_value=0)
        .reset_index()
    )

    summary = code_counts.copy()
    count_cols = [c for c in summary.columns if c != "Resident"]

    summary["Worked Shifts"] = summary[count_cols].apply(
        lambda row: sum(row[code] for code in count_cols if hours_map.get(code, 0) > 0),
        axis=1,
    )

    summary["Total Hours"] = summary[count_cols].apply(
        lambda row: sum(row[code] * hours_map.get(code, 0) for code in count_cols),
        axis=1,
    )

    summary["Avg Hours / Week"] = summary["Total Hours"] / max(num_weeks, 1)
    summary["Avg Shifts / Week"] = summary["Worked Shifts"] / max(num_weeks, 1)
    summary["OFF Days / 7 Days"] = summary["OFF"] / max(num_weeks, 1) if "OFF" in summary.columns else 0
    summary["POST Days / 7 Days"] = summary["POST"] / max(num_weeks, 1) if "POST" in summary.columns else 0

    return summary


def sheet_ref(sheet_name: str) -> str:
    return f"'{sheet_name.replace(chr(39), chr(39) + chr(39))}'"


def count_formula_for_rows(
    sheet_name: str,
    schedule_rows: list[int],
    criteria_ref: str,
) -> str:
    pieces = [
        f"COUNTIF({sheet_ref(sheet_name)}!$B${row}:$H${row},{criteria_ref})"
        for row in schedule_rows
    ]
    return f"=SUM({','.join(pieces)})"


def make_excel(
    schedule_df: pd.DataFrame,
    shift_defs: pd.DataFrame,
    residents: list[str],
    num_weeks: int,
) -> bytes:
    shift_codes = shift_defs["Code"].tolist()
    n_residents = len(residents)
    block_height = n_residents + 1  # one header row plus resident rows

    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    summary_ws = wb.create_sheet("Summary")
    defs_ws = wb.create_sheet("Shift_Definitions")

    title_fill = PatternFill("solid", fgColor="1F4E78")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    week_fill = PatternFill("solid", fgColor="E2F0D9")
    thin_gray = Side(style="thin", color="B7B7B7")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    # Shift definitions sheet
    defs_ws["A1"] = "Shift Code"
    defs_ws["B1"] = "Hours"

    for cell in defs_ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = title_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    for i, row in shift_defs.iterrows():
        excel_row = i + 2
        defs_ws.cell(excel_row, 1).value = row["Code"]
        defs_ws.cell(excel_row, 2).value = float(row["Hours"])
        defs_ws.cell(excel_row, 1).border = border
        defs_ws.cell(excel_row, 2).border = border

    defs_ws.column_dimensions["A"].width = 16
    defs_ws.column_dimensions["B"].width = 10

    # Schedule sheet
    ws["A1"] = "Resident Schedule"
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=14)

    ws["A2"] = (
        f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}. "
        "Edit shift cells directly; formulas update in Excel."
    )
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=14)

    schedule_headers = [
        "Resident",
        *DAYS,
        "D Count",
        "N Count",
        "Worked Shifts",
        "Hours",
        "OFF Days",
        "POST Days",
    ]

    dv_formula = f"={sheet_ref('Shift_Definitions')}!$A$2:$A${len(shift_codes) + 1}"
    shift_validation = DataValidation(type="list", formula1=dv_formula, allow_blank=False)

    for week in range(1, num_weeks + 1):
        header_row = 3 + (week - 1) * block_height
        data_start = header_row + 1

        for col_idx, header in enumerate(schedule_headers, start=1):
            cell = ws.cell(header_row, col_idx)
            cell.value = header if col_idx != 1 else f"Week {week}"
            cell.font = Font(bold=True)
            cell.fill = week_fill if col_idx == 1 else header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

        for resident_idx, resident in enumerate(residents):
            row_num = data_start + resident_idx
            row_data = schedule_df[
                (schedule_df["Week"] == week) & (schedule_df["Resident"] == resident)
            ]

            ws.cell(row_num, 1).value = resident
            ws.cell(row_num, 1).font = Font(bold=True)
            ws.cell(row_num, 1).border = border

            for day_idx, day in enumerate(DAYS, start=2):
                val = clean_code(row_data.iloc[0][day]) if not row_data.empty else "OFF"
                ws.cell(row_num, day_idx).value = val
                ws.cell(row_num, day_idx).alignment = Alignment(horizontal="center")
                ws.cell(row_num, day_idx).border = border

            row_range = f"$B{row_num}:$H{row_num}"

            ws.cell(row_num, 9).value = f'=COUNTIF({row_range},"D")'
            ws.cell(row_num, 10).value = f'=COUNTIF({row_range},"N")'

            worked_terms = []
            hour_terms = []

            for def_idx in range(len(shift_codes)):
                def_row = def_idx + 2
                code_cell = f"{sheet_ref('Shift_Definitions')}!$A${def_row}"
                hour_cell = f"{sheet_ref('Shift_Definitions')}!$B${def_row}"

                worked_terms.append(f"COUNTIF({row_range},{code_cell})*({hour_cell}>0)")
                hour_terms.append(f"COUNTIF({row_range},{code_cell})*{hour_cell}")

            ws.cell(row_num, 11).value = f"=SUM({','.join(worked_terms)})"
            ws.cell(row_num, 12).value = f"=SUM({','.join(hour_terms)})"
            ws.cell(row_num, 13).value = f'=COUNTIF({row_range},"OFF")'
            ws.cell(row_num, 14).value = f'=COUNTIF({row_range},"POST")'

            for col_idx in range(9, 15):
                ws.cell(row_num, col_idx).alignment = Alignment(horizontal="center")
                ws.cell(row_num, col_idx).border = border

        ws.add_data_validation(shift_validation)
        shift_validation.add(f"B{data_start}:H{data_start + n_residents - 1}")

    ws.freeze_panes = "B4"

    for col, width in {
        "A": 18,
        "B": 10,
        "C": 10,
        "D": 10,
        "E": 10,
        "F": 10,
        "G": 10,
        "H": 10,
        "I": 11,
        "J": 11,
        "K": 15,
        "L": 10,
        "M": 10,
        "N": 10,
    }.items():
        ws.column_dimensions[col].width = width

    # Summary sheet
    summary_headers = [
        "Resident",
        *shift_codes,
        "Worked Shifts",
        "Total Hours",
        "Avg Hours / Week",
        "Avg Shifts / Week",
        "OFF Days / 7 Days",
        "POST Days / 7 Days",
    ]

    summary_ws["A1"] = "Shift Counter and Weekly Averages"
    summary_ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    summary_ws["A1"].fill = title_fill
    summary_ws["A1"].alignment = Alignment(horizontal="center")
    summary_ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(summary_headers))

    header_row = 3

    for col_idx, header in enumerate(summary_headers, start=1):
        cell = summary_ws.cell(header_row, col_idx)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    code_start_col = 2
    code_end_col = code_start_col + len(shift_codes) - 1
    worked_col = code_end_col + 1
    hours_col = worked_col + 1
    avg_hours_col = hours_col + 1
    avg_shifts_col = avg_hours_col + 1
    off_col = avg_shifts_col + 1
    post_col = off_col + 1

    for resident_idx, resident in enumerate(residents):
        summary_row = header_row + 1 + resident_idx

        summary_ws.cell(summary_row, 1).value = resident
        summary_ws.cell(summary_row, 1).font = Font(bold=True)
        summary_ws.cell(summary_row, 1).border = border

        schedule_rows = [
            4 + resident_idx + (week_idx * block_height)
            for week_idx in range(num_weeks)
        ]

        for code_idx, _code in enumerate(shift_codes):
            col_idx = code_start_col + code_idx
            col_letter = get_column_letter(col_idx)
            criteria_ref = f"{col_letter}${header_row}"

            summary_ws.cell(summary_row, col_idx).value = count_formula_for_rows(
                "Schedule",
                schedule_rows,
                criteria_ref,
            )
            summary_ws.cell(summary_row, col_idx).border = border
            summary_ws.cell(summary_row, col_idx).alignment = Alignment(horizontal="center")

        count_terms_for_worked = []
        count_terms_for_hours = []

        for code_idx in range(len(shift_codes)):
            count_col = get_column_letter(code_start_col + code_idx)
            def_row = code_idx + 2
            hour_cell = f"{sheet_ref('Shift_Definitions')}!$B${def_row}"
            count_cell = f"{count_col}{summary_row}"

            count_terms_for_worked.append(f"{count_cell}*({hour_cell}>0)")
            count_terms_for_hours.append(f"{count_cell}*{hour_cell}")

        worked_letter = get_column_letter(worked_col)
        hours_letter = get_column_letter(hours_col)

        summary_ws.cell(summary_row, worked_col).value = f"=SUM({','.join(count_terms_for_worked)})"
        summary_ws.cell(summary_row, hours_col).value = f"=SUM({','.join(count_terms_for_hours)})"
        summary_ws.cell(summary_row, avg_hours_col).value = f"={hours_letter}{summary_row}/{num_weeks}"
        summary_ws.cell(summary_row, avg_shifts_col).value = f"={worked_letter}{summary_row}/{num_weeks}"

        code_header_range = (
            f"${get_column_letter(code_start_col)}${header_row}:"
            f"${get_column_letter(code_end_col)}${header_row}"
        )
        code_count_range = (
            f"{get_column_letter(code_start_col)}{summary_row}:"
            f"{get_column_letter(code_end_col)}{summary_row}"
        )

        summary_ws.cell(summary_row, off_col).value = (
            f'=IFERROR(INDEX({code_count_range},1,MATCH("OFF",{code_header_range},0))/{num_weeks},0)'
        )
        summary_ws.cell(summary_row, post_col).value = (
            f'=IFERROR(INDEX({code_count_range},1,MATCH("POST",{code_header_range},0))/{num_weeks},0)'
        )

        for col_idx in range(worked_col, post_col + 1):
            summary_ws.cell(summary_row, col_idx).border = border
            summary_ws.cell(summary_row, col_idx).alignment = Alignment(horizontal="center")

            if col_idx in [avg_hours_col, avg_shifts_col, off_col, post_col]:
                summary_ws.cell(summary_row, col_idx).number_format = "0.00"

    summary_ws.freeze_panes = "B4"

    for col_idx in range(1, len(summary_headers) + 1):
        col_letter = get_column_letter(col_idx)
        summary_ws.column_dimensions[col_letter].width = max(
            12,
            min(22, len(str(summary_headers[col_idx - 1])) + 2),
        )

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return output.getvalue()


def main() -> None:
    st.set_page_config(page_title="Resident Schedule Builder", layout="wide")

    st.title("Resident Schedule Builder")
    st.caption(
        "Build D/N/OFF/POST schedules, count shifts, calculate hours, "
        "and export to Excel with formulas."
    )

    st.sidebar.header("Setup")

    num_residents = int(
        st.sidebar.number_input(
            "Number of residents",
            min_value=1,
            max_value=20,
            value=3,
            step=1,
        )
    )

    num_weeks = int(
        st.sidebar.number_input(
            "Number of weeks",
            min_value=1,
            max_value=26,
            value=5,
            step=1,
        )
    )

    residents = make_resident_names(num_residents)
    shift_defs = make_shift_definitions()
    shift_codes = shift_defs["Code"].tolist()
    default_code = "OFF" if "OFF" in shift_codes else shift_codes[0]

    signature = (tuple(residents), num_weeks, tuple(shift_codes))

    if "schedule_df" not in st.session_state:
        st.session_state.schedule_df = build_blank_schedule(
            residents,
            num_weeks,
            default_code,
        )
        st.session_state.schedule_signature = signature

    if st.session_state.get("schedule_signature") != signature:
        st.session_state.schedule_df = reconcile_schedule(
            st.session_state.schedule_df,
            residents,
            num_weeks,
            default_code,
        )
        st.session_state.schedule_signature = signature

        # Clear stale editor state when rows/options change.
        for key in list(st.session_state.keys()):
            if key.startswith("schedule_editor_week_") or key == "schedule_editor":
                del st.session_state[key]

    #st.subheader("Schedule")
    #st.write("Each week is shown separately, matching the Excel layout. No internal table scroll.")

    column_config = {
        day: st.column_config.SelectboxColumn(
            day,
            options=shift_codes,
            required=True,
        )
        for day in DAYS
    }

    disabled_cols = [
        "Resident",
        "D Count",
        "N Count",
        "Worked Shifts",
        "Hours",
        "OFF Days",
        "POST Days",
    ]

    for week in range(1, num_weeks + 1):
        st.markdown(f"### Week {week}")

        week_df = st.session_state.schedule_df[
            st.session_state.schedule_df["Week"] == week
        ][["Resident", *DAYS]].reset_index(drop=True)

        week_display_df = add_week_calculations(week_df, shift_defs)

        st.data_editor(
            week_display_df,
            hide_index=True,
            use_container_width=True,
            num_rows="fixed",
            disabled=disabled_cols,
            column_config=column_config,
            key=f"schedule_editor_week_{week}",
            on_change=sync_week_editor,
            args=(week,),
            height=38 * (len(residents) + 1) + 8,
        )

    schedule_df = st.session_state.schedule_df.copy()

    for day in DAYS:
        schedule_df[day] = schedule_df[day].map(clean_code)

    st.subheader("Overall Shift Counter")
    summary_df = summarize_schedule(schedule_df, shift_defs, residents, num_weeks)
    st.dataframe(summary_df, use_container_width=True, hide_index=True)

    st.subheader("Export")
    excel_bytes = make_excel(schedule_df, shift_defs, residents, num_weeks)

    st.download_button(
        label="Download Excel schedule with formulas",
        data=excel_bytes,
        file_name="resident_schedule_with_formulas.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    with st.expander("Excel formula logic"):
        st.markdown(
            """
            The exported workbook has:

            - **Schedule** tab: editable schedule cells with dropdowns.
            - **Summary** tab: formulas that count each shift code across repeated weekly rows.
            - **Shift_Definitions** tab: shift codes and hours.

            The summary formulas are generated automatically and follow the same idea as:

            `=SUM(COUNTIF($B$4:$H$4,"D"),COUNTIF($B$8:$H$8,"D"),COUNTIF($B$12:$H$12,"D"))`
            """
        )


if __name__ == "__main__":
    main()
